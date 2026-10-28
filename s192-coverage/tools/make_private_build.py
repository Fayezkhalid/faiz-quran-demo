#!/usr/bin/env python3
"""Builds a private, ready-to-use copy of the S192 coverage app with data embedded.

Usage: make_private_build.py ROSTER.xls MANPOWER.xlsx OUT.html [--store S192] [--exclude ID,ID] [--names names_ar.json] [--artifact] [--no-seed] [--sync-ops ops.json] [--logo logo.png]
  --names     JSON file {employee_id: "الاسم بالعربي"} (keep it outside the repo)
  --no-seed   do not embed the data (cloud-synced artifact reads it from the db)
  --sync-ops  also write the initial db documents as JSON [{path, data}]
  --artifact  strip the <html>/<head>/<body> wrapper (for publishing as an Artifact page)
The output contains employee names and phone numbers: never commit or publish it.
"""
import sys, re, json, time
from html.parser import HTMLParser
import openpyxl

import argparse
ap = argparse.ArgumentParser()
ap.add_argument("roster"); ap.add_argument("manpower"); ap.add_argument("out")
ap.add_argument("--store", default="S192"); ap.add_argument("--exclude", default="")
ap.add_argument("--names"); ap.add_argument("--artifact", action="store_true")
ap.add_argument("--no-seed", action="store_true"); ap.add_argument("--logo"); ap.add_argument("--sync-ops")
a = ap.parse_args()
roster_path, man_path, out_path, store = a.roster, a.manpower, a.out, a.store
exclude = {x.strip() for x in a.exclude.split(",") if x.strip()}
names_ar = json.load(open(a.names, encoding="utf-8")) if a.names else {}

def group_of(pos, dept):
    p = (pos or "").upper()
    if p == "CASHIER": return "cashier"
    if p.startswith("CUSTOMER SERVICE"): return "cs"
    if "PRODUCE" in p or re.search(r"_PROD$", dept or "", re.I): return "produce"
    return "other"

rows, st = [], {"cur": None, "cell": None}
class P(HTMLParser):
    def handle_starttag(s, t, a):
        if t == "tr": st["cur"] = []
        if t in ("td", "th"): st["cell"] = ""
    def handle_endtag(s, t):
        if t in ("td", "th") and st["cur"] is not None and st["cell"] is not None:
            st["cur"].append(re.sub(r"[\s\xa0]+", " ", st["cell"]).strip()); st["cell"] = None
        if t == "tr" and st["cur"] is not None: rows.append(st["cur"]); st["cur"] = None
    def handle_data(s, d):
        if st["cell"] is not None: st["cell"] += d
P().feed(open(roster_path, "rb").read().decode("cp1256", errors="replace"))
hdr = next(r for r in rows if len(r) >= 4 and r[0].lower() == "employee code" and re.search(r"\d+/\d+/\d{4}", r[2]))
dates = []
for h in hdr[2:]:
    d, m, y = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", h).groups()
    dates.append(f"{y}-{int(m):02d}-{int(d):02d}")
ros = {r[0]: {"name": r[1], "shifts": r[2:]} for r in rows if len(r) == len(hdr) and r[0].isdigit()}

wb = openpyxl.load_workbook(man_path, read_only=True, data_only=True)
ws = wb[next(n for n in wb.sheetnames if "employees" in n.lower())]
it = ws.iter_rows(values_only=True)
H = None
for r in it:
    if any(str(c).strip().lower() == "nationality" for c in r if c): H = [str(c or "").strip().lower() for c in r]; break
ix = {k: H.index(v) for k, v in dict(id="employee", nat="nationality", name="employee name", mob="mobile", cc="cost center", pos="position", store="store").items()}
staff, known = {}, []
for r in it:
    if str(r[ix["store"]]).strip() == store: known.append(re.sub(r"\D", "", str(r[ix["id"]])))
    if str(r[ix["store"]]).strip() != store or "saudi" not in str(r[ix["nat"]]).lower(): continue
    pos = str(r[ix["pos"]] or "").strip()
    if pos.upper() == "STORE MANAGER": continue
    i = re.sub(r"\D", "", str(r[ix["id"]]))
    if i in exclude: continue
    staff[i] = dict(id=i, name=re.sub(r"\s+", " ", str(r[ix["name"]])).strip(), pos=pos, dept=str(r[ix["cc"]] or ""),
                    phone=re.sub(r"\D", "", str(r[ix["mob"]] or "")), grp=group_of(pos, r[ix["cc"]]))
    if i in names_ar: staff[i]["ar"] = names_ar[i]

seed = {"stamp": str(int(time.time())), "store": store, "staff": staff, "known": known, "roster": {"dates": dates, "rows": ros}}
if a.sync_ops:
    ops = [{"path": f"staff/{i}", "data": {"rec": rec, "removed": False}} for i, rec in staff.items()]
    ops.append({"path": "meta/known", "data": {"ids": known}})
    ops.append({"path": "meta/roster", "data": {"dates": dates, "rows": ros, "at": int(time.time()) * 1000}})
    json.dump(ops, open(a.sync_ops, "w", encoding="utf-8"), ensure_ascii=False)
src = open(__file__.replace("tools/make_private_build.py", "index.html"), encoding="utf-8").read()
if a.logo:
    import base64
    seed["logo"] = "data:image/png;base64," + base64.b64encode(open(a.logo, "rb").read()).decode()
if a.no_seed: seed = None
src = src.replace('<script id="seed" type="application/json">null</script>',
                  '<script id="seed" type="application/json">' + json.dumps(seed, ensure_ascii=False).replace("</", "<\\/") + "</script>")
if a.artifact:
    src = re.sub(r"<!doctype html>\s*<html[^>]*>\s*<head>", "", src, flags=re.I)
    src = src.replace('<meta charset="utf-8">\n', "").replace('<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n', "")
    src = src.replace("</head>\n<body>", "").replace("</body>\n</html>", "")
    src = src.replace("<title>", '<script>document.documentElement.dir="rtl";document.documentElement.lang="ar"</script>\n<title>', 1)
    src = src.replace("padding:calc(18px + env(safe-area-inset-top,0px)) 0 18px", "padding:18px 0")
open(out_path, "w", encoding="utf-8").write(src)
print(f"{len(staff)} Saudi staff, {len(ros)} roster rows, {dates[0]}..{dates[-1]} -> {out_path}")
